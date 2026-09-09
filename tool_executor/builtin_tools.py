"""Deterministic, reviewed file and web tools for the platform executor.

The executor never installs packages or runs user-provided code. Every exposed
operation is implemented in this immutable module and validated before use.
"""

from __future__ import annotations

import csv
import html
import ipaddress
import json
import mimetypes
import re
import shutil
import socket
import stat
import subprocess
import tarfile
import tempfile
import urllib.parse
import urllib.request
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any, ClassVar

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter


class BuiltinToolError(ValueError):
    """A platform file operation could not be completed."""


MAX_WEB_TEXT_BYTES = 2 * 1024 * 1024
MAX_WEB_DOWNLOAD_BYTES = 5 * 1024 * 1024
MAX_ARCHIVE_FILES = 20
MAX_ARCHIVE_FILE_BYTES = 5 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = 40_000_000

_ALLOWED_ACTIONS = {
    "spreadsheet": {"inspect", "create", "edit", "convert"},
    "document": {"inspect", "create", "edit", "convert"},
    "presentation": {"inspect", "create", "edit", "convert"},
    "pdf": {"inspect", "create", "merge", "split", "extract_pages", "convert"},
    "text": {"inspect", "create", "edit", "convert"},
    "web": {"search", "fetch", "download"},
    "image": {"inspect", "convert", "resize", "crop", "compress", "ocr"},
    "archive": {"list", "extract", "create"},
}


def _assert_safe_zip_package(path: Path) -> None:
    with path.open("rb") as handle:
        signature = handle.read(4)
    if signature != b"PK\x03\x04":
        return
    try:
        with zipfile.ZipFile(path) as package:
            infos = package.infolist()
            if len(infos) > 10_000:
                raise BuiltinToolError("文件包包含的项目过多，已拒绝处理")
            expanded = 0
            for info in infos:
                member = PurePosixPath(info.filename.replace("\\", "/"))
                if member.is_absolute() or ".." in member.parts:
                    raise BuiltinToolError("文件包包含不安全路径，已拒绝处理")
                expanded += max(0, int(info.file_size))
                if expanded > 500 * 1024 * 1024:
                    raise BuiltinToolError("文件解压后的体积超过 500MB，已拒绝处理")
                if info.compress_size and info.file_size / info.compress_size > 1_000:
                    raise BuiltinToolError("文件压缩率异常，可能是压缩炸弹，已拒绝处理")
    except zipfile.BadZipFile as exc:
        raise BuiltinToolError("ZIP 或 Office 文件包已损坏") from exc


def validate_builtin_request(
    tool_kind: str, action: str, inputs: list[Path], params: dict
) -> None:
    """Fail closed at the executor boundary even if a caller bypasses model schemas."""

    if action not in _ALLOWED_ACTIONS.get(tool_kind, set()):
        raise BuiltinToolError(f"不支持的文件操作：{tool_kind}.{action}")
    for field in ("sheets", "slides", "operations", "pages"):
        if field in params and isinstance(params[field], str):
            raise BuiltinToolError(f"参数 {field} 必须是真实数组，不能是 JSON 字符串")
    required_input = action in {
        "inspect",
        "edit",
        "convert",
        "merge",
        "split",
        "extract_pages",
        "ocr",
        "resize",
        "crop",
        "compress",
        "list",
        "extract",
    }
    if required_input and not inputs:
        raise BuiltinToolError("该文件操作缺少输入文件")
    for source in inputs:
        _assert_safe_zip_package(source)
    if tool_kind in {"document", "presentation", "text"} and action == "edit":
        mode = str(
            params.get("mode") or ("replace" if params.get("replace") else "append")
        )
        if tool_kind in {"document", "presentation"} and mode != "append":
            raise BuiltinToolError("为避免丢失原格式，Office 文件只支持明确的追加编辑")


def validate_output_file(path: Path) -> tuple[str, str]:
    """Reopen every output and verify its real format before it leaves the sandbox."""

    if not path.is_file() or path.stat().st_size <= 0:
        raise BuiltinToolError(f"输出文件为空：{path.name}")
    suffix = path.suffix.lower().lstrip(".")
    mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    if suffix in {"xlsx", "xlsm", "xltx", "xltm"}:
        try:
            book = load_workbook(
                path, read_only=True, keep_vba=suffix in {"xlsm", "xltm"}
            )
            if not book.sheetnames:
                raise ValueError("workbook has no sheets")
            book.close()
        except Exception as exc:
            raise BuiltinToolError(f"Excel 无法重新打开：{path.name}") from exc
    elif suffix in {"docx", "docm", "dotx", "dotm"}:
        try:
            Document(path)
        except Exception as exc:
            raise BuiltinToolError(f"Word 无法重新打开：{path.name}") from exc
    elif suffix in {"pptx", "pptm", "ppsx", "ppsm", "potx", "potm"}:
        try:
            presentation = Presentation(path)
            if len(presentation.slides) < 1:
                raise ValueError("presentation has no slides")
        except Exception as exc:
            raise BuiltinToolError(f"PowerPoint 无法重新打开：{path.name}") from exc
    elif suffix in {"ods", "odt", "odp", "ots", "ott", "otp"}:
        try:
            with zipfile.ZipFile(path) as package:
                names = set(package.namelist())
            if "mimetype" not in names or "content.xml" not in names:
                raise ValueError("invalid OpenDocument package")
        except Exception as exc:
            raise BuiltinToolError(f"OpenDocument 无法重新打开：{path.name}") from exc
    elif suffix in {"xls", "doc", "ppt", "pps", "pot"}:
        if not path.read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1"):
            raise BuiltinToolError(f"旧版 Office 文件头无效：{path.name}")
    elif suffix == "pdf":
        try:
            reader = PdfReader(path, strict=True)
            if len(reader.pages) < 1:
                raise ValueError("PDF has no pages")
            rendered = fitz.open(path)
            rendered[0].get_pixmap(matrix=fitz.Matrix(0.25, 0.25), alpha=False)
            rendered.close()
        except Exception as exc:
            raise BuiltinToolError(f"PDF 无法重新打开或渲染：{path.name}") from exc
    elif suffix in {"txt", "md", "markdown", "csv", "tsv"}:
        try:
            path.read_text(encoding="utf-8-sig", errors="strict")
        except UnicodeDecodeError as exc:
            raise BuiltinToolError(f"文本文件不是有效 UTF-8：{path.name}") from exc
    elif suffix in {"png", "jpg", "jpeg", "webp", "tiff", "bmp"}:
        try:
            with Image.open(path) as image:
                image.verify()
        except Exception as exc:
            raise BuiltinToolError(f"图片无法重新打开：{path.name}") from exc
    elif suffix == "zip":
        try:
            with zipfile.ZipFile(path) as archive:
                if archive.testzip() is not None:
                    raise ValueError("corrupt archive member")
        except Exception as exc:
            raise BuiltinToolError(f"ZIP 无法重新打开：{path.name}") from exc
    return suffix, mime


def validate_output_semantics(
    path: Path,
    tool_kind: str,
    action: str,
    params: dict,
) -> None:
    """Check deterministic structure promised by create operations."""

    if action != "create":
        return
    suffix = path.suffix.lower()
    if tool_kind == "spreadsheet" and suffix in {".xlsx", ".xlsm"}:
        book = load_workbook(path, read_only=True)
        try:
            expected = [
                str(item.get("name") or f"Sheet{index + 1}")[:31]
                for index, item in enumerate(params.get("sheets") or [])
            ]
            if expected and book.sheetnames != expected:
                raise BuiltinToolError("Excel 工作表数量或名称与生成请求不一致")
        finally:
            book.close()
    elif tool_kind == "document" and suffix == ".docx":
        document = Document(path)
        source_text = str(params.get("markdown") or params.get("content") or "").strip()
        delivered_text = "\n".join(
            paragraph.text for paragraph in document.paragraphs
        ).strip()
        if source_text and not delivered_text and not document.tables:
            raise BuiltinToolError("Word 主要段落或表格未正确写入")
    elif tool_kind == "presentation" and suffix == ".pptx":
        presentation = Presentation(path)
        expected_slides = len(params.get("slides") or [])
        if len(presentation.slides) != expected_slides:
            raise BuiltinToolError("PowerPoint 幻灯片数量与生成请求不一致")


def _value(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    # Excel may contain library-specific scalar values. Inspection must never
    # turn an otherwise readable workbook into a Runner 500 response.
    return json.dumps(value, ensure_ascii=False, default=str)


def _safe_output_name(value: str | None, default: str, suffix: str) -> str:
    name = Path((value or default).replace("\\", "/")).name
    if not name.lower().endswith(suffix):
        name += suffix
    return name


def _libreoffice_convert(source: Path, output_dir: Path, target: str) -> Path:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise BuiltinToolError("文件兼容转换服务当前不可用")
    profile = Path(tempfile.mkdtemp(prefix="builtin-lo-"))
    try:
        result = subprocess.run(
            [
                executable,
                "--headless",
                f"-env:UserInstallation={profile.as_uri()}",
                "--convert-to",
                target,
                "--outdir",
                str(output_dir),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
        if result.returncode:
            detail = (result.stderr or result.stdout or "未知错误")[-1000:]
            raise BuiltinToolError(f"文件格式转换失败：{detail}")
        candidates = sorted(
            output_dir.glob(f"{source.stem}.*"), key=lambda item: item.stat().st_mtime
        )
        if not candidates:
            raise BuiltinToolError("文件格式转换失败：没有生成目标文件")
        return candidates[-1]
    except subprocess.TimeoutExpired as exc:
        raise BuiltinToolError("文件格式转换超时，请缩小文件后重试") from exc
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def _load_tabular(path: Path):
    if path.suffix.lower() in {".csv", ".tsv"}:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        rows = list(
            csv.reader(
                path.read_text(encoding="utf-8-sig").splitlines(), delimiter=delimiter
            )
        )
        book = Workbook()
        sheet = book.active
        sheet.title = path.stem[:31] or "Sheet1"
        for row in rows:
            sheet.append(row)
        return book
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        converted_dir = Path(tempfile.mkdtemp(prefix="builtin-sheet-"))
        converted = _libreoffice_convert(path, converted_dir, "xlsx")
        book = load_workbook(converted, data_only=False)
        shutil.rmtree(converted_dir, ignore_errors=True)
        return book
    return load_workbook(
        path,
        data_only=False,
        keep_vba=path.suffix.lower() in {".xlsm", ".xltm"},
    )


def _style_sheet(sheet) -> None:
    if sheet.max_row:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.alignment = Alignment(vertical="center")
    for column in range(1, min(sheet.max_column, 100) + 1):
        width = 10
        for row in range(1, min(sheet.max_row, 200) + 1):
            value = sheet.cell(row, column).value
            if value is not None:
                width = max(width, min(len(str(value)) + 2, 42))
        sheet.column_dimensions[get_column_letter(column)].width = width


def _write_sheets(book: Workbook, sheets: list[dict]) -> None:
    while book.worksheets:
        book.remove(book.worksheets[0])
    for index, spec in enumerate(sheets or [{"name": "Sheet1", "rows": []}]):
        sheet = book.create_sheet(str(spec.get("name") or f"Sheet{index + 1}")[:31])
        for row in spec.get("rows") or []:
            sheet.append([_value(value) for value in row])
        _style_sheet(sheet)


def _spreadsheet(
    action: str, inputs: list[Path], params: dict, output_dir: Path
) -> dict:
    if action == "inspect":
        if not inputs:
            raise BuiltinToolError("spreadsheet inspect requires one input file")
        book = _load_tabular(inputs[0])
        requested_sheet = str(params.get("sheet") or "").strip()
        if requested_sheet and requested_sheet not in book.sheetnames:
            raise BuiltinToolError(
                f"spreadsheet sheet does not exist: {requested_sheet}"
            )
        selected_sheets = (
            [book[requested_sheet]] if requested_sheet else list(book.worksheets)
        )
        max_cols = min(max(int(params.get("max_columns", 30)), 1), 100)
        cell_range = str(params.get("range") or "").strip().upper()
        offset = max(int(params.get("offset", 0)), 0)
        limit = min(max(int(params.get("limit", params.get("max_rows", 50))), 1), 1000)
        sheets = []
        any_more = False
        next_offsets: list[int] = []
        for sheet in selected_sheets:
            if cell_range:
                try:
                    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
                except ValueError as exc:
                    raise BuiltinToolError(
                        "spreadsheet range must be an A1 range such as A2:F200"
                    ) from exc
                if max_col - min_col + 1 > 100 or max_row - min_row + 1 > 1000:
                    raise BuiltinToolError(
                        "spreadsheet range exceeds 100 columns or 1000 rows"
                    )
                effective_max_col = min(max_col, min_col + max_cols - 1)
                row_start = min_row
                row_end = min(max_row, sheet.max_row)
            else:
                min_col = 1
                effective_max_col = min(sheet.max_column, max_cols)
                row_start = offset + 1
                row_end = min(sheet.max_row, offset + limit)
            rows = [
                [_value(cell.value) for cell in row]
                for row in sheet.iter_rows(
                    min_row=row_start,
                    max_row=max(row_start - 1, row_end),
                    min_col=min_col,
                    max_col=max(min_col, effective_max_col),
                )
            ]
            has_more = not cell_range and row_end < sheet.max_row
            next_offset = row_end if has_more else None
            any_more = any_more or has_more
            if next_offset is not None:
                next_offsets.append(next_offset)
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": rows,
                    "total_rows": sheet.max_row,
                    "total_columns": sheet.max_column,
                    "offset": (row_start - 1),
                    "limit": (row_end - row_start + 1) if row_end >= row_start else 0,
                    "range": cell_range or None,
                    "has_more": has_more,
                    "next_offset": next_offset,
                }
            )
        return {
            "summary": {
                "kind": "spreadsheet",
                "sheets": sheets,
                "has_more": any_more,
                "next_offset": min(next_offsets) if next_offsets else None,
            },
            "outputs": [],
        }

    if action == "create":
        target = str(params.get("target_format") or "xlsx").lower().lstrip(".")
        sheets_spec = params.get("sheets") or [
            {"name": "Sheet1", "rows": params.get("rows") or []}
        ]
        if target in {"csv", "tsv"}:
            if len(sheets_spec) != 1:
                raise BuiltinToolError("CSV/TSV 只能包含一个工作表，请只指定一个工作表")
            suffix = f".{target}"
            output = output_dir / _safe_output_name(
                params.get("output_name"), "workbook" + suffix, suffix
            )
            delimiter = "\t" if target == "tsv" else ","
            with output.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, delimiter=delimiter)
                writer.writerows(sheets_spec[0].get("rows") or [])
            return {"summary": f"已创建 {target.upper()} 文件", "outputs": [output]}
        if target not in {"xlsx", "xls"}:
            raise BuiltinToolError(f"不支持创建此表格格式：{target}")
        book = Workbook()
        _write_sheets(book, sheets_spec)
        if target == "xls":
            source = output_dir / "source.xlsx"
            book.save(source)
            produced = _libreoffice_convert(source, output_dir, "xls")
            source.unlink(missing_ok=True)
            final = output_dir / _safe_output_name(
                params.get("output_name"), "workbook.xls", ".xls"
            )
            if produced != final:
                produced.replace(final)
            return {"summary": "已创建兼容格式 XLS 文件", "outputs": [final]}
    elif action == "edit":
        source_suffix = inputs[0].suffix.lower()
        if source_suffix not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            raise BuiltinToolError(
                "表格编辑需要现代 Excel 文件；旧格式或文本表格请先转换为 XLSX"
            )
        if source_suffix in {".xlsm", ".xltm"}:
            book = load_workbook(inputs[0], data_only=False, keep_vba=False)
        else:
            book = _load_tabular(inputs[0])
        for operation in params.get("operations") or []:
            op = operation.get("type")
            sheet_name = str(operation.get("sheet") or book.sheetnames[0])
            if sheet_name not in book.sheetnames:
                book.create_sheet(sheet_name[:31])
            sheet = book[sheet_name[:31]]
            if op == "set_cell":
                sheet[str(operation.get("cell") or "A1")] = _value(
                    operation.get("value")
                )
            elif op == "append_rows":
                for row in operation.get("rows") or []:
                    sheet.append([_value(value) for value in row])
            elif op == "replace_rows":
                sheet.delete_rows(1, sheet.max_row)
                for row in operation.get("rows") or []:
                    sheet.append([_value(value) for value in row])
            elif op == "rename_sheet":
                sheet.title = str(operation.get("new_name") or sheet.title)[:31]
            else:
                raise BuiltinToolError(f"Unsupported spreadsheet edit operation: {op}")
    elif action == "convert":
        if not inputs:
            raise BuiltinToolError("spreadsheet convert requires one input file")
        target = str(params.get("target_format") or "xlsx").lower().lstrip(".")
        if target not in {"xlsx", "csv", "tsv", "ods", "pdf", "xls"}:
            raise BuiltinToolError(f"不支持转换为此表格格式：{target}")
        if target in {"ods", "pdf", "xls"}:
            produced = _libreoffice_convert(inputs[0], output_dir, target)
            requested = _safe_output_name(
                params.get("output_name"), produced.name, f".{target}"
            )
            final = output_dir / requested
            if produced != final:
                produced.replace(final)
            return {"summary": f"converted spreadsheet to {target}", "outputs": [final]}
        book = _load_tabular(inputs[0])
        if target in {"csv", "tsv"}:
            requested_sheet = str(params.get("sheet") or "").strip()
            if len(book.sheetnames) > 1 and not requested_sheet:
                raise BuiltinToolError(
                    "多工作表文件转 CSV/TSV 时必须指定 sheet，避免静默丢失数据"
                )
            if requested_sheet and requested_sheet not in book.sheetnames:
                raise BuiltinToolError(f"指定的工作表不存在：{requested_sheet}")
            suffix = f".{target}"
            output = output_dir / _safe_output_name(
                params.get("output_name"), inputs[0].stem, suffix
            )
            delimiter = "\t" if target == "tsv" else ","
            with output.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, delimiter=delimiter)
                selected = requested_sheet or book.sheetnames[0]
                for row in book[selected].iter_rows(values_only=True):
                    writer.writerow(list(row))
            return {
                "summary": f"converted spreadsheet to {target}",
                "outputs": [output],
            }
    else:
        raise BuiltinToolError(f"Unsupported spreadsheet action: {action}")

    if action == "edit":
        source_suffix = inputs[0].suffix.lower()
        if source_suffix in {".xlsm", ".xltm", ".xltx"}:
            source_suffix = ".xlsx"
        output = output_dir / _safe_output_name(
            params.get("output_name"),
            inputs[0].stem + source_suffix,
            source_suffix,
        )
    else:
        output = output_dir / _safe_output_name(
            params.get("output_name"), "workbook.xlsx", ".xlsx"
        )
    book.save(output)
    summary = "表格文件已创建" if action == "create" else "表格文件已更新"
    if action == "edit" and inputs[0].suffix.lower() in {".xlsm", ".xltm"}:
        summary += "；原文件未覆盖，已生成不含宏的 XLSX 副本"
    return {"summary": summary, "outputs": [output]}


def _append_markdown(document: Document, markdown: str) -> None:
    lines = markdown.replace("\r\n", "\n").split("\n")
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line[:3].rstrip(".").isdigit() and ". " in line[:5]:
            document.add_paragraph(line.split(". ", 1)[1], style="List Number")
        elif (
            line.startswith("|")
            and index + 1 < len(lines)
            and not lines[index + 1]
            .replace("|", "")
            .replace("-", "")
            .replace(":", "")
            .strip()
        ):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip("|").split("|")]
                if not all(set(cell) <= {"-", ":", " "} for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                table = document.add_table(
                    rows=len(rows), cols=max(len(row) for row in rows)
                )
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for column_index, value in enumerate(row):
                        table.cell(row_index, column_index).text = value
            continue
        elif line.strip():
            document.add_paragraph(line)
        index += 1


def _document(action: str, inputs: list[Path], params: dict, output_dir: Path) -> dict:
    if action == "inspect":
        if not inputs:
            raise BuiltinToolError("document inspect requires one input file")
        source = inputs[0]
        temp: Path | None = None
        if source.suffix.lower() != ".docx":
            temp = Path(tempfile.mkdtemp(prefix="builtin-doc-"))
            source = _libreoffice_convert(source, temp, "docx")
        document = Document(source)
        summary = {
            "kind": "document",
            "paragraphs": [paragraph.text for paragraph in document.paragraphs[:500]],
            "tables": [
                [[cell.text for cell in row.cells] for row in table.rows]
                for table in document.tables[:20]
            ],
        }
        if temp:
            shutil.rmtree(temp, ignore_errors=True)
        return {"summary": summary, "outputs": []}
    if action == "create":
        document = Document()
        _append_markdown(
            document, str(params.get("markdown") or params.get("content") or "")
        )
        target = str(params.get("target_format") or "docx").lower().lstrip(".")
        if target not in {"docx", "doc"}:
            raise BuiltinToolError(f"不支持创建此 Word 格式：{target}")
        if target == "doc":
            source = output_dir / "source.docx"
            document.save(source)
            produced = _libreoffice_convert(source, output_dir, "doc")
            source.unlink(missing_ok=True)
            final = output_dir / _safe_output_name(
                params.get("output_name"), "document.doc", ".doc"
            )
            if produced != final:
                produced.replace(final)
            return {"summary": "已创建兼容格式 DOC 文件", "outputs": [final]}
    elif action == "edit":
        if not inputs:
            raise BuiltinToolError("document edit requires one DOCX input file")
        if inputs[0].suffix.lower() != ".docx":
            raise BuiltinToolError(
                "document edit currently requires DOCX; convert legacy files first"
            )
        if params.get("replace"):
            document = Document()
        else:
            document = Document(inputs[0])
        _append_markdown(
            document, str(params.get("markdown") or params.get("content") or "")
        )
    elif action == "convert":
        if not inputs:
            raise BuiltinToolError("document convert requires one input file")
        target = str(params.get("target_format") or "pdf").lower().lstrip(".")
        if target not in {"pdf", "docx", "doc", "odt", "rtf"}:
            raise BuiltinToolError(f"不支持转换为此 Word 格式：{target}")
        produced = _libreoffice_convert(inputs[0], output_dir, target)
        final = output_dir / _safe_output_name(
            params.get("output_name"), produced.name, f".{target}"
        )
        if produced != final:
            produced.replace(final)
        return {"summary": f"converted document to {target}", "outputs": [final]}
    else:
        raise BuiltinToolError(f"Unsupported document action: {action}")
    output = output_dir / _safe_output_name(
        params.get("output_name"), "document.docx", ".docx"
    )
    document.save(output)
    return {
        "summary": "Word 文件已创建" if action == "create" else "Word 文件已更新",
        "outputs": [output],
    }


def _add_slides(presentation: Presentation, slides: list[dict]) -> None:
    for spec in slides:
        layout = (
            presentation.slide_layouts[1]
            if len(presentation.slide_layouts) > 1
            else presentation.slide_layouts[0]
        )
        slide = presentation.slides.add_slide(layout)
        if slide.shapes.title:
            slide.shapes.title.text = str(spec.get("title") or "")
        body = "\n".join(str(value) for value in (spec.get("bullets") or []))
        placeholders = [
            shape for shape in slide.placeholders if shape != slide.shapes.title
        ]
        if placeholders and hasattr(placeholders[0], "text_frame"):
            placeholders[0].text_frame.text = body
        elif body:
            box = slide.shapes.add_textbox(
                Inches(1), Inches(1.8), Inches(8), Inches(4.5)
            )
            box.text_frame.text = body
        if spec.get("notes") and slide.notes_slide.notes_text_frame:
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])


def _presentation(
    action: str, inputs: list[Path], params: dict, output_dir: Path
) -> dict:
    if action == "inspect":
        if not inputs:
            raise BuiltinToolError("presentation inspect requires one input file")
        source = inputs[0]
        temp: Path | None = None
        if source.suffix.lower() != ".pptx":
            temp = Path(tempfile.mkdtemp(prefix="builtin-ppt-"))
            source = _libreoffice_convert(source, temp, "pptx")
        presentation = Presentation(source)
        slides = []
        for index, slide in enumerate(presentation.slides):
            slides.append(
                {
                    "number": index + 1,
                    "texts": [
                        shape.text
                        for shape in slide.shapes
                        if hasattr(shape, "text") and shape.text
                    ],
                }
            )
        if temp:
            shutil.rmtree(temp, ignore_errors=True)
        return {"summary": {"kind": "presentation", "slides": slides}, "outputs": []}
    if action == "create":
        presentation = Presentation()
        _add_slides(presentation, params.get("slides") or [])
        target = str(params.get("target_format") or "pptx").lower().lstrip(".")
        if target not in {"pptx", "ppt"}:
            raise BuiltinToolError(f"不支持创建此 PowerPoint 格式：{target}")
        if target == "ppt":
            source = output_dir / "source.pptx"
            presentation.save(source)
            produced = _libreoffice_convert(source, output_dir, "ppt")
            source.unlink(missing_ok=True)
            final = output_dir / _safe_output_name(
                params.get("output_name"), "presentation.ppt", ".ppt"
            )
            if produced != final:
                produced.replace(final)
            return {"summary": "已创建兼容格式 PPT 文件", "outputs": [final]}
    elif action == "edit":
        if not inputs or inputs[0].suffix.lower() != ".pptx":
            raise BuiltinToolError("presentation edit requires one PPTX input file")
        presentation = Presentation(inputs[0])
        _add_slides(presentation, params.get("slides") or [])
    elif action == "convert":
        if not inputs:
            raise BuiltinToolError("presentation convert requires one input file")
        target = str(params.get("target_format") or "pdf").lower().lstrip(".")
        if target not in {"pdf", "pptx", "ppt", "odp"}:
            raise BuiltinToolError(f"不支持转换为此 PowerPoint 格式：{target}")
        produced = _libreoffice_convert(inputs[0], output_dir, target)
        final = output_dir / _safe_output_name(
            params.get("output_name"), produced.name, f".{target}"
        )
        if produced != final:
            produced.replace(final)
        return {"summary": f"converted presentation to {target}", "outputs": [final]}
    else:
        raise BuiltinToolError(f"Unsupported presentation action: {action}")
    output = output_dir / _safe_output_name(
        params.get("output_name"), "presentation.pptx", ".pptx"
    )
    presentation.save(output)
    return {
        "summary": "PowerPoint 已创建" if action == "create" else "PowerPoint 已更新",
        "outputs": [output],
    }


def _pdf_page_texts(
    source: Path, max_pages: int, language: str
) -> tuple[list[str], bool]:
    reader = PdfReader(source)
    pages = [(page.extract_text() or "").strip() for page in reader.pages[:max_pages]]
    if any(pages):
        return pages, False
    try:
        import pytesseract
    except ImportError as exc:  # pragma: no cover - production image includes it
        raise BuiltinToolError("扫描 PDF 需要 OCR，但 OCR 组件当前不可用") from exc
    try:
        pages = [
            pytesseract.image_to_string(image, lang=language).strip()
            for image in _ocr_images(source, max_pages)
        ]
    except Exception as exc:
        raise BuiltinToolError(
            "扫描 PDF 的 OCR 识别失败，请检查语言包或文件质量"
        ) from exc
    return pages, True


def _pdf(action: str, inputs: list[Path], params: dict, output_dir: Path) -> dict:
    if action == "inspect":
        reader = PdfReader(inputs[0])
        max_pages = min(max(int(params.get("max_pages", 20)), 1), 100)
        language = str(params.get("ocr_language") or "chi_sim+eng")
        pages, ocr_used = _pdf_page_texts(inputs[0], max_pages, language)
        return {
            "summary": {
                "kind": "pdf",
                "page_count": len(reader.pages),
                "pages": [text[:20_000] for text in pages],
                "ocr_used": ocr_used,
            },
            "outputs": [],
        }
    if action == "create":
        document = Document()
        _append_markdown(
            document, str(params.get("markdown") or params.get("content") or "")
        )
        temp_docx = output_dir / "source.docx"
        document.save(temp_docx)
        produced = _libreoffice_convert(temp_docx, output_dir, "pdf")
        temp_docx.unlink(missing_ok=True)
        final = output_dir / _safe_output_name(
            params.get("output_name"), "document.pdf", ".pdf"
        )
        if produced != final:
            produced.replace(final)
        return {"summary": "PDF 已创建", "outputs": [final]}
    if action in {"merge", "split", "extract_pages"}:
        writer = PdfWriter()
        if action == "merge":
            for source in inputs:
                for page in PdfReader(source).pages:
                    writer.add_page(page)
            output = output_dir / _safe_output_name(
                params.get("output_name"), "merged.pdf", ".pdf"
            )
            with output.open("wb") as handle:
                writer.write(handle)
            return {"summary": "PDF 合并完成", "outputs": [output]}
        if action == "extract_pages":
            reader = PdfReader(inputs[0])
            pages = params.get("pages") or []
            for number in pages:
                index = int(number) - 1
                if index < 0 or index >= len(reader.pages):
                    raise BuiltinToolError(f"PDF 第 {number} 页超出范围")
                writer.add_page(reader.pages[index])
            output = output_dir / _safe_output_name(
                params.get("output_name"), "extracted.pdf", ".pdf"
            )
            with output.open("wb") as handle:
                writer.write(handle)
            return {"summary": "PDF 页面抽取完成", "outputs": [output]}
        reader = PdfReader(inputs[0])
        requested_pages = params.get("pages") or list(range(1, len(reader.pages) + 1))
        outputs: list[Path] = []
        requested_name = Path(str(params.get("output_name") or inputs[0].stem)).stem
        for number in requested_pages:
            index = int(number) - 1
            if index < 0 or index >= len(reader.pages):
                raise BuiltinToolError(f"PDF 第 {number} 页超出范围")
            page_writer = PdfWriter()
            page_writer.add_page(reader.pages[index])
            output = output_dir / f"{requested_name}-第{number}页.pdf"
            with output.open("wb") as handle:
                page_writer.write(handle)
            outputs.append(output)
        return {"summary": f"PDF 已拆分为 {len(outputs)} 个文件", "outputs": outputs}
    if action == "convert":
        target = str(params.get("target_format") or "txt").lower().lstrip(".")
        if target != "txt":
            raise BuiltinToolError("PDF 当前只支持转换为 TXT")
        max_pages = min(len(PdfReader(inputs[0]).pages), 100)
        pages, ocr_used = _pdf_page_texts(
            inputs[0],
            max_pages,
            str(params.get("ocr_language") or "chi_sim+eng"),
        )
        text = "\n\n".join(pages)
        output = output_dir / _safe_output_name(
            params.get("output_name"), inputs[0].stem, ".txt"
        )
        output.write_text(text, encoding="utf-8")
        return {
            "summary": {"message": "PDF 已转换为文本", "ocr_used": ocr_used},
            "outputs": [output],
        }
    raise BuiltinToolError(f"不支持的 PDF 操作：{action}")


def _markdown_to_plain_text(content: str) -> str:
    """Conservative CommonMark-to-text conversion without executing embedded HTML."""
    value = re.sub(r"```[^\n]*\n(.*?)```", r"\1", content, flags=re.DOTALL)
    value = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"(?m)^\s{0,3}#{1,6}\s+", "", value)
    value = re.sub(r"(?m)^\s*(?:[-+*]|\d+[.)])\s+", "", value)
    value = re.sub(r"(?<!\\)[*_~`]", "", value)
    return value


def _text(action: str, inputs: list[Path], params: dict, output_dir: Path) -> dict:
    if action == "inspect":
        try:
            content = inputs[0].read_text(encoding="utf-8-sig", errors="strict")
        except UnicodeDecodeError as exc:
            raise BuiltinToolError("文本文件不是有效 UTF-8，请先转换编码") from exc
        offset = max(int(params.get("offset") or 0), 0)
        limit = min(max(int(params.get("limit") or 100_000), 1), 100_000)
        end = min(offset + limit, len(content))
        return {
            "summary": {
                "kind": "text",
                "characters": len(content),
                "content": content[offset:end],
                "offset": offset,
                "limit": end - offset,
                "has_more": end < len(content),
                "next_offset": end if end < len(content) else None,
            },
            "outputs": [],
        }
    suffix = (
        ".md"
        if str(params.get("format") or "").lower() in {"md", "markdown"}
        else ".txt"
    )
    if action == "create":
        content = str(params.get("content") or "")
    elif action == "edit":
        if not inputs:
            raise BuiltinToolError("text edit requires one input file")
        try:
            original = inputs[0].read_text(encoding="utf-8-sig", errors="strict")
        except UnicodeDecodeError as exc:
            raise BuiltinToolError("文本文件不是有效 UTF-8，请先转换编码") from exc
        if str(params.get("mode") or "append") == "replace":
            content = str(params.get("content") or "")
        else:
            content = original + str(params.get("content") or "")
        suffix = (
            inputs[0].suffix.lower()
            if inputs[0].suffix.lower() in {".txt", ".md"}
            else suffix
        )
    elif action == "convert":
        if not inputs:
            raise BuiltinToolError("text convert requires one input file")
        try:
            content = inputs[0].read_text(encoding="utf-8-sig", errors="strict")
        except UnicodeDecodeError as exc:
            raise BuiltinToolError("文本文件不是有效 UTF-8，请先转换编码") from exc
        target = str(params.get("target_format") or "txt").lower().lstrip(".")
        if target not in {"txt", "md"}:
            raise BuiltinToolError(f"不支持转换为此文本格式：{target}")
        if target == "txt" and inputs[0].suffix.lower() in {".md", ".markdown"}:
            content = _markdown_to_plain_text(content)
        suffix = f".{target}"
    else:
        raise BuiltinToolError(f"不支持的文本操作：{action}")
    output = output_dir / _safe_output_name(
        params.get("output_name"), "document" + suffix, suffix
    )
    output.write_text(content, encoding="utf-8")
    return {"summary": "文本文件已处理", "outputs": [output]}


class _ReadableHTMLParser(HTMLParser):
    """Extract a useful title and readable text without running page scripts."""

    _SKIP_TAGS: ClassVar[set[str]] = {"script", "style", "noscript", "svg"}
    _BLOCK_TAGS: ClassVar[set[str]] = {
        "article",
        "aside",
        "blockquote",
        "br",
        "div",
        "footer",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "header",
        "li",
        "main",
        "nav",
        "p",
        "section",
        "table",
        "td",
        "th",
        "tr",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip_depth = 0
        self._in_title = False
        self.title_parts: list[str] = []
        self.text_parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        tag = tag.lower()
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
        if self._skip_depth:
            return
        if tag == "title":
            self._in_title = True
        if tag in self._BLOCK_TAGS:
            self.text_parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in self._SKIP_TAGS and self._skip_depth:
            self._skip_depth -= 1
            return
        if self._skip_depth:
            return
        if tag == "title":
            self._in_title = False
        if tag in self._BLOCK_TAGS:
            self.text_parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        value = data.strip()
        if not value:
            return
        if self._in_title:
            self.title_parts.append(value)
        self.text_parts.append(value + " ")

    def result(self) -> tuple[str, str]:
        title = " ".join(self.title_parts).strip()
        text = "".join(self.text_parts)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n\s*\n+", "\n\n", text)
        return title, text.strip()


class _DuckDuckGoParser(HTMLParser):
    """Parse the stable DuckDuckGo HTML result page without a browser."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.results: list[dict[str, str]] = []
        self._capture: str | None = None
        self._parts: list[str] = []
        self._href = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = {key: value or "" for key, value in attrs}
        classes = set(values.get("class", "").split())
        if tag == "a" and "result__a" in classes:
            self._capture = "title"
            self._parts = []
            self._href = values.get("href", "")
        elif tag in {"a", "div"} and "result__snippet" in classes:
            self._capture = "snippet"
            self._parts = []

    def handle_endtag(self, tag: str) -> None:
        if self._capture == "title" and tag == "a":
            title = " ".join(self._parts).strip()
            if title and self._href:
                parsed = urllib.parse.urlparse(self._href)
                query = urllib.parse.parse_qs(parsed.query)
                href = query.get("uddg", [self._href])[0]
                self.results.append({"title": title, "url": href, "snippet": ""})
            self._capture = None
        elif self._capture == "snippet" and tag in {"a", "div"}:
            if self.results:
                self.results[-1]["snippet"] = " ".join(self._parts).strip()
            self._capture = None

    def handle_data(self, data: str) -> None:
        if self._capture and data.strip():
            self._parts.append(data.strip())


class _BingParser(HTMLParser):
    """Parse Bing's non-JavaScript result markup as a search fallback."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.results: list[dict[str, str]] = []
        self._current: dict[str, str] | None = None
        self._in_heading = False
        self._capture: str | None = None
        self._parts: list[str] = []

    def _finish_result(self) -> None:
        if self._current and self._current.get("title") and self._current.get("url"):
            self.results.append(self._current)
        self._current = None
        self._in_heading = False
        self._capture = None
        self._parts = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = {key: value or "" for key, value in attrs}
        classes = set(values.get("class", "").split())
        if tag == "li" and "b_algo" in classes:
            self._finish_result()
            self._current = {"title": "", "url": "", "snippet": ""}
            return
        if self._current is None:
            return
        if tag == "h2":
            self._in_heading = True
        elif tag == "a" and self._in_heading:
            self._capture = "title"
            self._parts = []
            self._current["url"] = values.get("href", "")
        elif tag == "p":
            self._capture = "snippet"
            self._parts = []

    def handle_endtag(self, tag: str) -> None:
        if self._current is None:
            return
        if tag == "a" and self._capture == "title":
            self._current["title"] = " ".join(self._parts).strip()
            self._capture = None
        elif tag == "p" and self._capture == "snippet":
            self._current["snippet"] = " ".join(self._parts).strip()
            self._capture = None
        elif tag == "h2":
            self._in_heading = False
        elif tag == "li":
            self._finish_result()

    def handle_data(self, data: str) -> None:
        if self._capture and data.strip():
            self._parts.append(data.strip())

    def close(self) -> None:
        super().close()
        self._finish_result()


def _validate_public_url(url: str) -> urllib.parse.ParseResult:
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise BuiltinToolError("web_tool only accepts absolute HTTP/HTTPS URLs")
    hostname = parsed.hostname.rstrip(".").lower()
    if hostname == "localhost" or hostname.endswith(".local"):
        raise BuiltinToolError("Local or private network URLs are not allowed")
    try:
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
        addresses = {item[4][0] for item in socket.getaddrinfo(hostname, port)}
    except ValueError as exc:
        raise BuiltinToolError("Invalid URL port") from exc
    except socket.gaierror as exc:
        raise BuiltinToolError(f"Unable to resolve host: {hostname}") from exc
    for value in addresses:
        address = ipaddress.ip_address(value.split("%", 1)[0])
        if (
            address.is_private
            or address.is_loopback
            or address.is_link_local
            or address.is_multicast
            or address.is_reserved
            or address.is_unspecified
        ):
            raise BuiltinToolError("Local or private network URLs are not allowed")
    return parsed


class _SafeRedirectHandler(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        _validate_public_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def _http_get(url: str, *, limit: int) -> tuple[str, bytes, str, str]:
    _validate_public_url(url)
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "AI-Platform-WebTool/1.0 (+https://ai-platform.staging.zhuojianai.com)",
            "Accept": "text/html,application/json,text/plain,*/*;q=0.5",
        },
    )
    opener = urllib.request.build_opener(_SafeRedirectHandler())
    try:
        with opener.open(request, timeout=20) as response:
            raw = response.read(limit + 1)
            if len(raw) > limit:
                raise BuiltinToolError(
                    f"Remote response exceeds {limit // (1024 * 1024)}MB"
                )
            content_type = (
                response.headers.get_content_type() or "application/octet-stream"
            )
            disposition = response.headers.get("Content-Disposition", "")
            return response.geturl(), raw, content_type, disposition
    except BuiltinToolError:
        raise
    except Exception as exc:
        raise BuiltinToolError(f"Web request failed: {exc}") from exc


def _decode_web_text(raw: bytes, content_type: str) -> str:
    del content_type
    for encoding in ("utf-8", "gb18030", "big5"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _download_name(url: str, disposition: str, content_type: str) -> str:
    match = re.search(
        r"filename\*?=(?:UTF-8''|\")?([^\";]+)", disposition, flags=re.IGNORECASE
    )
    if match:
        name = urllib.parse.unquote(match.group(1).strip())
    else:
        name = urllib.parse.unquote(PurePosixPath(urllib.parse.urlparse(url).path).name)
    if not name:
        name = "download" + (mimetypes.guess_extension(content_type) or ".bin")
    return Path(name.replace("\\", "/")).name


def _web(action: str, inputs: list[Path], params: dict, output_dir: Path) -> dict:
    del inputs
    if action == "search":
        query = str(params.get("query") or "").strip()
        if not query:
            raise BuiltinToolError("web search requires query")
        max_results = min(max(int(params.get("max_results", 5)), 1), 10)
        url = "https://html.duckduckgo.com/html/?" + urllib.parse.urlencode(
            {"q": query}
        )
        _, raw, content_type, _ = _http_get(url, limit=MAX_WEB_TEXT_BYTES)
        parser = _DuckDuckGoParser()
        parser.feed(_decode_web_text(raw, content_type))
        results = parser.results[:max_results]
        if not results:
            fallback_url = "https://www.bing.com/search?" + urllib.parse.urlencode(
                {"q": query, "count": max_results}
            )
            _, raw, content_type, _ = _http_get(fallback_url, limit=MAX_WEB_TEXT_BYTES)
            fallback = _BingParser()
            fallback.feed(_decode_web_text(raw, content_type))
            fallback.close()
            results = fallback.results[:max_results]
        if not results:
            raise BuiltinToolError("Search providers returned no parseable results")
        return {"summary": {"query": query, "results": results}, "outputs": []}
    if action == "fetch":
        url = str(params.get("url") or "").strip()
        final_url, raw, content_type, _ = _http_get(url, limit=MAX_WEB_TEXT_BYTES)
        text = _decode_web_text(raw, content_type)
        title = ""
        if (
            content_type in {"text/html", "application/xhtml+xml"}
            or "<html" in text[:500].lower()
        ):
            parser = _ReadableHTMLParser()
            parser.feed(text)
            title, text = parser.result()
        max_chars = min(max(int(params.get("max_chars", 50_000)), 1_000), 100_000)
        return {
            "summary": {
                "url": final_url,
                "title": html.unescape(title),
                "content_type": content_type,
                "content": text[:max_chars],
                "truncated": len(text) > max_chars,
            },
            "outputs": [],
        }
    if action == "download":
        url = str(params.get("url") or "").strip()
        final_url, raw, content_type, disposition = _http_get(
            url,
            limit=MAX_WEB_DOWNLOAD_BYTES,
        )
        name = _safe_output_name(
            params.get("output_name"),
            _download_name(final_url, disposition, content_type),
            "",
        )
        output = output_dir / name
        output.write_bytes(raw)
        return {
            "summary": {
                "url": final_url,
                "content_type": content_type,
                "size": len(raw),
            },
            "outputs": [output],
        }
    raise BuiltinToolError(f"Unsupported web action: {action}")


def _image_output_suffix(value: str | None, fallback: str = "png") -> tuple[str, str]:
    normalized = str(value or fallback).lower().lstrip(".")
    aliases = {"jpg": "jpeg", "tif": "tiff"}
    normalized = aliases.get(normalized, normalized)
    if normalized not in {"png", "jpeg", "webp", "tiff", "bmp"}:
        raise BuiltinToolError(f"Unsupported image format: {normalized}")
    suffix = ".jpg" if normalized == "jpeg" else f".{normalized}"
    return normalized.upper(), suffix


def _save_image(
    image: Image.Image, output: Path, image_format: str, quality: int
) -> None:
    if image_format == "JPEG" and image.mode not in {"RGB", "L"}:
        background = Image.new("RGB", image.size, "white")
        if "A" in image.getbands():
            background.paste(image, mask=image.getchannel("A"))
        else:
            background.paste(image.convert("RGB"))
        image = background
    kwargs: dict[str, Any] = {"format": image_format}
    if image_format in {"JPEG", "WEBP"}:
        kwargs.update({"quality": quality, "optimize": True})
    elif image_format == "PNG":
        kwargs["optimize"] = True
    image.save(output, **kwargs)


def _ocr_images(source: Path, max_pages: int) -> list[Image.Image]:
    if source.suffix.lower() == ".pdf":
        document = fitz.open(source)
        images: list[Image.Image] = []
        try:
            for page in document[:max_pages]:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                images.append(
                    Image.frombytes(
                        "RGB", (pixmap.width, pixmap.height), pixmap.samples
                    )
                )
        finally:
            document.close()
        return images
    with Image.open(source) as image:
        return [ImageOps.exif_transpose(image).convert("RGB")]


def _image(action: str, inputs: list[Path], params: dict, output_dir: Path) -> dict:
    if not inputs:
        raise BuiltinToolError(f"image {action} requires an input file")
    source = inputs[0]
    if action == "inspect":
        if source.suffix.lower() == ".pdf":
            document = fitz.open(source)
            try:
                return {
                    "summary": {
                        "kind": "scanned_document",
                        "pages": document.page_count,
                    },
                    "outputs": [],
                }
            finally:
                document.close()
        with Image.open(source) as image:
            return {
                "summary": {
                    "kind": "image",
                    "format": image.format,
                    "width": image.width,
                    "height": image.height,
                    "mode": image.mode,
                    "frames": getattr(image, "n_frames", 1),
                },
                "outputs": [],
            }
    if action == "ocr":
        try:
            import pytesseract
        except ImportError as exc:  # pragma: no cover - production image includes it
            raise BuiltinToolError("OCR runtime is unavailable") from exc
        language = str(params.get("language") or "chi_sim+eng")
        max_pages = min(max(int(params.get("max_pages", 10)), 1), 20)
        texts = [
            pytesseract.image_to_string(image, lang=language)
            for image in _ocr_images(source, max_pages)
        ]
        content = "\n\n".join(text.strip() for text in texts).strip()
        outputs: list[Path] = []
        if params.get("output_name"):
            output = output_dir / _safe_output_name(
                params.get("output_name"), "ocr.txt", ".txt"
            )
            output.write_text(content, encoding="utf-8")
            outputs.append(output)
        return {
            "summary": {
                "language": language,
                "pages": len(texts),
                "content": content[:100_000],
            },
            "outputs": outputs,
        }
    if source.suffix.lower() == ".pdf":
        raise BuiltinToolError(
            f"image {action} does not accept PDF; use ocr for scanned PDFs"
        )
    with Image.open(source) as opened:
        image = ImageOps.exif_transpose(opened).copy()
    image_format, suffix = _image_output_suffix(
        params.get("target_format"),
        (source.suffix.lower().lstrip(".") or "png"),
    )
    if action == "resize":
        width = int(params.get("width") or image.width)
        height = int(params.get("height") or image.height)
        if width < 1 or height < 1 or width * height > 40_000_000:
            raise BuiltinToolError("Invalid or oversized image dimensions")
        if params.get("keep_aspect", True):
            image.thumbnail((width, height), Image.Resampling.LANCZOS)
        else:
            image = image.resize((width, height), Image.Resampling.LANCZOS)
    elif action == "crop":
        box = params.get("box") or []
        if not isinstance(box, list) or len(box) != 4:
            raise BuiltinToolError("image crop requires box=[left, top, right, bottom]")
        coordinates = tuple(int(value) for value in box)
        if (
            coordinates[0] < 0
            or coordinates[1] < 0
            or coordinates[2] > image.width
            or coordinates[3] > image.height
        ):
            raise BuiltinToolError("Crop box is outside the image")
        if coordinates[2] <= coordinates[0] or coordinates[3] <= coordinates[1]:
            raise BuiltinToolError("Crop box has no area")
        image = image.crop(coordinates)
    elif action not in {"convert", "compress"}:
        raise BuiltinToolError(f"Unsupported image action: {action}")
    quality = min(max(int(params.get("quality", 85)), 1), 100)
    output = output_dir / _safe_output_name(
        params.get("output_name"), "image" + suffix, suffix
    )
    _save_image(image, output, image_format, quality)
    return {
        "summary": {"action": action, "width": image.width, "height": image.height},
        "outputs": [output],
    }


def _safe_archive_path(name: str) -> PurePosixPath:
    normalized = PurePosixPath(name.replace("\\", "/").lstrip("/"))
    if not normalized.parts or ".." in normalized.parts or normalized.is_absolute():
        raise BuiltinToolError(f"Unsafe archive path: {name}")
    return normalized


def _archive_kind(path: Path) -> str:
    name = path.name.lower()
    if name.endswith(".zip"):
        return "zip"
    if name.endswith((".tar.gz", ".tgz")):
        return "tar.gz"
    if name.endswith(".tar"):
        return "tar"
    raise BuiltinToolError("archive_tool supports ZIP, TAR and TAR.GZ files")


def _archive_list(source: Path) -> list[dict[str, Any]]:
    kind = _archive_kind(source)
    if kind == "zip":
        with zipfile.ZipFile(source) as archive:
            return [
                {"path": item.filename, "size": item.file_size, "is_dir": item.is_dir()}
                for item in archive.infolist()
            ]
    with tarfile.open(source, "r:gz" if kind == "tar.gz" else "r:") as archive:
        return [
            {"path": item.name, "size": item.size, "is_dir": item.isdir()}
            for item in archive.getmembers()
        ]


def _archive_extract(source: Path, output_dir: Path) -> list[Path]:
    kind = _archive_kind(source)
    outputs: list[Path] = []
    if kind == "zip":
        with zipfile.ZipFile(source) as archive:
            members = [item for item in archive.infolist() if not item.is_dir()]
            if len(members) > MAX_ARCHIVE_FILES:
                raise BuiltinToolError(
                    f"Archive contains more than {MAX_ARCHIVE_FILES} files"
                )
            for item in members:
                if item.flag_bits & 0x1:
                    raise BuiltinToolError("Encrypted ZIP files are not supported")
                if stat.S_IFMT(item.external_attr >> 16) == stat.S_IFLNK:
                    raise BuiltinToolError("Archive links are not supported")
                if item.file_size > MAX_ARCHIVE_FILE_BYTES:
                    raise BuiltinToolError(
                        f"Archive member {item.filename} exceeds 5MB"
                    )
                relative = _safe_archive_path(item.filename)
                target = output_dir.joinpath(*relative.parts)
                target.parent.mkdir(parents=True, exist_ok=True)
                with (
                    archive.open(item) as source_file,
                    target.open("wb") as target_file,
                ):
                    shutil.copyfileobj(source_file, target_file)
                outputs.append(target)
        return outputs
    with tarfile.open(source, "r:gz" if kind == "tar.gz" else "r:") as archive:
        members = [item for item in archive.getmembers() if item.isfile()]
        if len(members) > MAX_ARCHIVE_FILES:
            raise BuiltinToolError(
                f"Archive contains more than {MAX_ARCHIVE_FILES} files"
            )
        for item in members:
            if item.issym() or item.islnk():
                raise BuiltinToolError("Archive links are not supported")
            if item.size > MAX_ARCHIVE_FILE_BYTES:
                raise BuiltinToolError(f"Archive member {item.name} exceeds 5MB")
            relative = _safe_archive_path(item.name)
            target = output_dir.joinpath(*relative.parts)
            target.parent.mkdir(parents=True, exist_ok=True)
            source_file = archive.extractfile(item)
            if source_file is None:
                continue
            with source_file, target.open("wb") as target_file:
                shutil.copyfileobj(source_file, target_file)
            outputs.append(target)
    return outputs


def _archive(action: str, inputs: list[Path], params: dict, output_dir: Path) -> dict:
    if action == "list":
        if len(inputs) != 1:
            raise BuiltinToolError("archive list requires one input file")
        members = _archive_list(inputs[0])
        return {
            "summary": {
                "format": _archive_kind(inputs[0]),
                "total": len(members),
                "members": members[:200],
                "truncated": len(members) > 200,
            },
            "outputs": [],
        }
    if action == "extract":
        if len(inputs) != 1:
            raise BuiltinToolError("archive extract requires one input file")
        outputs = _archive_extract(inputs[0], output_dir)
        return {"summary": {"extracted": len(outputs)}, "outputs": outputs}
    if action == "create":
        if not inputs:
            raise BuiltinToolError("archive create requires input files")
        kind = str(params.get("format") or "zip").lower()
        if kind not in {"zip", "tar", "tar.gz", "tgz"}:
            raise BuiltinToolError("archive format must be zip, tar or tar.gz")
        suffix = ".zip" if kind == "zip" else (".tar" if kind == "tar" else ".tar.gz")
        output = output_dir / _safe_output_name(
            params.get("output_name"), "archive" + suffix, suffix
        )
        if kind == "zip":
            with zipfile.ZipFile(
                output, "w", compression=zipfile.ZIP_DEFLATED
            ) as archive:
                for source in inputs:
                    archive.write(source, arcname=source.name)
        else:
            with tarfile.open(output, "w" if kind == "tar" else "w:gz") as archive:
                for source in inputs:
                    archive.add(source, arcname=source.name, recursive=False)
        return {
            "summary": {"created": output.name, "files": len(inputs)},
            "outputs": [output],
        }
    raise BuiltinToolError(f"Unsupported archive action: {action}")


def execute_builtin(
    tool_kind: str, action: str, inputs: list[Path], params: dict, output_dir: Path
) -> dict:
    handlers = {
        "spreadsheet": _spreadsheet,
        "document": _document,
        "presentation": _presentation,
        "pdf": _pdf,
        "text": _text,
        "web": _web,
        "image": _image,
        "archive": _archive,
    }
    handler = handlers.get(tool_kind)
    if handler is None:
        raise BuiltinToolError(f"Unsupported builtin tool kind: {tool_kind}")
    validate_builtin_request(tool_kind, action, inputs, params)
    result = handler(action, inputs, params, output_dir)
    mime_types: dict[str, str] = {}
    output_verification: dict[str, dict[str, str]] = {}
    for path in result.get("outputs") or []:
        detected_format, detected_mime = validate_output_file(path)
        validate_output_semantics(path, tool_kind, action, params)
        mime_types[path.name] = detected_mime
        output_verification[path.name] = {
            "detected_format": detected_format,
            "detected_mime_type": detected_mime,
        }
    result["mime_types"] = mime_types
    result["output_verification"] = output_verification
    return result
